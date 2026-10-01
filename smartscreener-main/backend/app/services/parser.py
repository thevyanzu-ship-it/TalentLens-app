import io
from pypdf import PdfReader
from docx import Document

class ResumeParserService:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        pdf_file = io.BytesIO(file_bytes)
        try:
            reader = PdfReader(pdf_file)
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF file: {str(e)}")

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        docx_file = io.BytesIO(file_bytes)
        try:
            doc = Document(docx_file)
            text_parts = []
            # Extract from paragraphs
            for para in doc.paragraphs:
                if para.text:
                    text_parts.append(para.text)
            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text]
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            return "\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file: {str(e)}")

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> str:
        ext = filename.split(".")[-1].lower()
        if ext == "pdf":
            text = cls.extract_text_from_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            text = cls.extract_text_from_docx(file_bytes)
        else:
            # Fallback for plain text files
            try:
                text = file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                raise ValueError(f"Unsupported file format: {ext}")
        
        # Clean text basic whitespace normalization
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines)
